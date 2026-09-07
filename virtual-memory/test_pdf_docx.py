#!/usr/bin/env python3
"""Validates splitter.py's content/structure separation (HANDBOOK.md
Experiment 17) against real, freshly-generated PDF and DOCX files -
not assumed, actually built and round-tripped.
"""
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from vstorage.splitter import split_bytes
from vstorage.system import VirtualStorage


def make_pdf(path: str) -> None:
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(path)
    c.drawString(100, 750, "Virtual Storage - PDF split test")
    for i in range(40):
        c.drawString(100, 700 - i * 15,
                     f"Line {i}: the quick brown fox jumps over the lazy dog.")
    c.save()


def make_docx(path: str) -> None:
    import docx
    d = docx.Document()
    d.add_heading("Virtual Storage - DOCX split test", level=1)
    for i in range(40):
        d.add_paragraph(f"Paragraph {i}: the quick brown fox jumps over the lazy dog.")
    d.save(path)


def main() -> None:
    tmpdir = "/dev/shm/vstorage_pdfdocx_test"
    os.makedirs(tmpdir, exist_ok=True)
    pdf_path = os.path.join(tmpdir, "sample.pdf")
    docx_path = os.path.join(tmpdir, "sample.docx")
    make_pdf(pdf_path)
    make_docx(docx_path)

    for label, path in [("PDF", pdf_path), ("DOCX", docx_path)]:
        raw = open(path, "rb").read()
        result = split_bytes(raw, path)
        print(f"\n{label}: {os.path.basename(path)}")
        print(f"  raw size:       {len(raw)} bytes")
        print(f"  content piece:  {len(result.content)} bytes "
              f"(reconstruct_from={result.reconstruct_from!r})")
        print(f"  structure piece:{len(result.structure)} bytes")
        print(f"  metadata piece: {result.metadata.decode()}")
        if result.content:
            preview = result.content[:80].decode("utf-8", errors="replace")
            print(f"  content preview: {preview!r}")

        vs = VirtualStorage()
        fid = vs.save_bytes(raw, path)
        full = vs.retrieve(fid, "full")
        print(f"  byte-perfect full reconstruction: {full == raw}")
        content_only = vs.retrieve(fid, "content")
        print(f"  content-only retrieval works: {content_only == result.content}")
        vs.forget(fid)

    for p in (pdf_path, docx_path):
        os.remove(p)
    os.rmdir(tmpdir)


if __name__ == "__main__":
    main()
