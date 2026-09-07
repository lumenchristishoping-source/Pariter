#!/usr/bin/env python3
"""Real-world compression check, per the user's request: use an actual
PDF they sent (not a synthetic one) and a bigger, more realistic DOCX.
Reports the two different ratios for PDF separately, since they answer
different questions:

  overall ratio   = raw file size / everything held (content+structure+
                     metadata, compressed) - this is "how much smaller is
                     what's sitting in RAM than the file on disk."
  content ratio    = raw / content piece alone (compressed) - "how much
                     smaller is just the extracted text."

For PDF/DOCX, reconstruct_from="structure", so the full original binary
is one of the three falling pieces. That binary is already a compressed
container (PDF streams, DOCX's own zip/deflate) - LZMA on top of already
-compressed bytes barely shrinks it (HANDBOOK.md section 9). So overall
ratio for PDF/DOCX will always be close to 1x - that's not a bug, it's
what "already compressed" means. The content ratio is where the real
interesting number is for these formats.
"""
import hashlib
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from vstorage.splitter import split_bytes
from vstorage.system import VirtualStorage

REAL_PDF = "/dev/shm/vstorage_real_test/real.pdf"


def make_real_docx(tmpdir: str) -> str:
    """A bigger, realistic DOCX - several headings, many paragraphs,
    a table - closer to a real document than the earlier 30-line memo."""
    import docx
    p = os.path.join(tmpdir, "realistic.docx")
    d = docx.Document()
    d.add_heading("Quarterly Operations Report", level=1)
    d.add_paragraph(
        "This report summarizes operational performance across all "
        "regional units for the current quarter, including staffing "
        "levels, throughput, and outstanding risk items that require "
        "leadership attention before the next review cycle."
    )
    for section in range(8):
        d.add_heading(f"Section {section + 1}: Regional Summary", level=2)
        for para in range(12):
            d.add_paragraph(
                f"In region {section + 1}, unit {para + 1} reported "
                "steady progress against targets this period. Staffing "
                "remained within budget, incident counts stayed below "
                "threshold, and no material risks were escalated. "
                "Follow-up actions are tracked in the shared log and "
                "will be reviewed again at the next checkpoint."
            )
    table = d.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = (
        "Region", "Headcount", "Incidents", "Status")
    for i in range(40):
        row = table.add_row().cells
        row[0].text = f"Region {i}"
        row[1].text = str(20 + i)
        row[2].text = str(i % 3)
        row[3].text = "OK" if i % 3 == 0 else "Watch"
    d.save(p)
    return p


def report(label: str, raw: bytes, split, vs: VirtualStorage, file_id: str) -> None:
    content_c = len(vs._compress(split.content))
    structure_c = len(vs._compress(split.structure))
    metadata_c = len(vs._compress(split.metadata))
    total_held = content_c + structure_c + metadata_c

    print(f"\n{label}")
    print(f"  raw file:              {len(raw):>9,} B")
    print(f"  content piece (text), compressed: {content_c:>9,} B "
          f"(raw content was {len(split.content):,} B)")
    print(f"  structure piece (full binary), compressed: {structure_c:>9,} B "
          f"(raw structure was {len(split.structure):,} B)")
    print(f"  metadata piece, compressed:        {metadata_c:>9,} B")
    print(f"  reconstruct_from:      {split.reconstruct_from!r}")
    print(f"  --")
    print(f"  overall ratio (raw / all 3 held):  {len(raw) / total_held:.2f}x")
    if split.content:
        print(f"  content-only ratio (raw / content held): "
              f"{len(raw) / content_c:.2f}x")
    else:
        print(f"  content-only ratio: n/a (no text extracted)")

    full = vs.retrieve(file_id, "full")
    match = hashlib.sha256(full).digest() == hashlib.sha256(raw).digest()
    print(f"  byte-perfect full retrieval: {match}")

    if split.content:
        text_back = vs.retrieve(file_id, "content")
        print(f"  content retrieval works: {len(text_back) == len(split.content)} "
              f"({len(text_back):,} B of extracted text)")


def main() -> None:
    vs = VirtualStorage()

    # --- Real PDF, sent by the user ---
    with open(REAL_PDF, "rb") as f:
        pdf_raw = f.read()
    pdf_split = split_bytes(pdf_raw, "real.pdf")
    pdf_id = vs.save_bytes(pdf_raw, "real.pdf")
    report("REAL PDF (user-supplied, FUTO_SM_Core_Portal1.PDF)",
           pdf_raw, pdf_split, vs, pdf_id)

    # --- Bigger, realistic DOCX (synthetic - no real one was supplied) ---
    tmpdir = "/dev/shm/vstorage_real_test"
    os.makedirs(tmpdir, exist_ok=True)
    docx_path = make_real_docx(tmpdir)
    with open(docx_path, "rb") as f:
        docx_raw = f.read()
    docx_split = split_bytes(docx_raw, "realistic.docx")
    docx_id = vs.save_bytes(docx_raw, "realistic.docx")
    report("DOCX (synthetic but realistic - 8 sections, 96 paragraphs, "
           "40-row table; no real DOCX was supplied, so this is generated)",
           docx_raw, docx_split, vs, docx_id)

    vs.collapse_all()
    os.remove(docx_path)


if __name__ == "__main__":
    main()
